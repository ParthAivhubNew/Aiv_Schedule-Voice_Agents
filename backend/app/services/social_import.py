"""Bulk import: many posts from one CSV or Excel file.

Each row is a post on one date and time, for one or more channels (one post per channel, as the
calendar keeps them). The file is checked first and every problem is shown per row; nothing is
saved until the user imports the rows that passed. A row with a topic but no text is written by
the AI (charged like any AI post). Imported posts go through approval like every other post.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time as dtime, timedelta
from types import SimpleNamespace
from typing import Any, Dict, List, Optional, Tuple

from app.services.social_publisher import TITLES, channel_issues, normalize_platform

MAX_ROWS = 500
MAX_FILE_BYTES = 5 * 1024 * 1024

# Column names people use, matched loosely (case, spaces and punctuation ignored).
COLUMNS = {
    "date": ("date", "day", "publishdate", "postdate", "scheduledate"),
    "time": ("time", "hour", "publishtime", "posttime"),
    "channels": ("channel", "channels", "network", "networks", "platform", "platforms"),
    "text": ("text", "caption", "copy", "post", "body", "message", "content", "posttext"),
    "title": ("title", "headline", "subject", "reddittitle"),
    "topic": ("topic", "idea", "brief", "about", "aitopic"),
    "media": ("media", "image", "images", "imageurl", "imageurls", "video", "videourl", "picture", "pictures", "mediaurl", "mediaurls"),
    "first_comment": ("firstcomment", "comment"),
    "alt_text": ("alt", "alttext", "imagedescription"),
    "subreddit": ("subreddit",),
}
TEMPLATE_HEADERS = ["date", "time", "channels", "title", "text", "topic", "media", "first comment", "alt text", "subreddit"]
MONTHS = {m: i + 1 for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"))}


def _key(header: Any) -> Optional[str]:
    h = re.sub(r"[^a-z0-9]", "", str(header or "").lower())
    for field, names in COLUMNS.items():
        if h in names:
            return field
    return None


def read_table(raw: bytes, filename: str) -> List[Dict[str, Any]]:
    """Rows as {field: value} for the columns we know. Raises ValueError in plain words."""
    name = (filename or "").lower()
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("That file is over 5 MB. Split it into smaller files.")
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ValueError("That Excel file could not be opened. Save it again as .xlsx or .csv.")
        grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        grid = [r for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValueError("Use a .csv or .xlsx file.")
    grid = [r for r in grid if any(str(c or "").strip() for c in r)]
    if not grid:
        raise ValueError("That file is empty.")
    keys = [_key(h) for h in grid[0]]
    if not any(keys):
        raise ValueError("The first row must name the columns, for example: date, time, channels, text. "
                         "Download the template to see them all.")
    if "date" not in keys or not ({"text", "topic"} & set(keys)):
        raise ValueError("The file needs a date column and a text (or topic) column.")
    rows = []
    for r in grid[1:]:
        row: Dict[str, Any] = {}
        for k, v in zip(keys, r):
            if k and v is not None and str(v).strip() != "":
                row[k] = v
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise ValueError(f"A file can hold up to {MAX_ROWS} posts; this has {len(rows)}. Split it into smaller files.")
    return rows


def parse_date(v: Any) -> Tuple[Optional[str], str]:
    """(YYYY-MM-DD, note). Day-first is assumed for 03/04/2026; the note says how it was read."""
    if isinstance(v, datetime):
        return v.date().isoformat(), ""
    if isinstance(v, date):
        return v.isoformat(), ""
    s = str(v or "").strip()
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:  # an Excel day number
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat(), ""
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return _iso(int(m[1]), int(m[2]), int(m[3])), ""
    m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})$", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y += 2000 if y < 100 else 0
        if a > 12:
            return _iso(y, b, a), ""
        if b > 12:
            return _iso(y, a, b), ""
        out = _iso(y, b, a)
        return out, (f"Read as {int(a)} {_month_name(b)} {y} (day first)." if out and a != b else "")
    m = re.match(r"^(\d{1,2})\s+([A-Za-z]{3})[a-z]*\.?,?\s+(\d{4})$", s) or None
    if m:
        return _iso(int(m[3]), MONTHS.get(m[2].lower(), 0), int(m[1])), ""
    m = re.match(r"^([A-Za-z]{3})[a-z]*\.?\s+(\d{1,2}),?\s+(\d{4})$", s)
    if m:
        return _iso(int(m[3]), MONTHS.get(m[1].lower(), 0), int(m[2])), ""
    return None, ""


def _month_name(n: int) -> str:
    return ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
            "November", "December")[n - 1]


def _iso(y: int, m: int, d: int) -> Optional[str]:
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return None


def parse_time(v: Any) -> Optional[str]:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (datetime, dtime)):
        return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, float) and 0 <= v < 1:  # Excel time as a fraction of a day
        mins = round(v * 1440)
        return f"{mins // 60:02d}:{mins % 60:02d}"
    s = str(v).strip().lower().replace(".", ":")
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?(?::\d{2})?\s*(am|pm)?$", s)
    if not m:
        return None
    h, mi = int(m[1]), int(m[2] or 0)
    if m[3] == "pm" and h < 12:
        h += 12
    if m[3] == "am" and h == 12:
        h = 0
    return f"{h:02d}:{mi:02d}" if h < 24 and mi < 60 else None


def split_list(v: Any) -> List[str]:
    return [p.strip() for p in re.split(r"[,;|\n]+", str(v or "")) if p.strip()]


def check_rows(rows: List[Dict[str, Any]], now_local: datetime, default_time: str = "09:00") -> List[Dict[str, Any]]:
    """Each row as the import will save it, with problems (block the row) and notes (it imports)."""
    out = []
    for n, r in enumerate(rows, start=2):  # row 1 holds the column names
        issues: List[str] = []
        notes: List[str] = []
        day, how = parse_date(r.get("date"))
        if how:
            notes.append(how)
        if not day:
            issues.append("Add a date, for example 2026-11-03." if not r.get("date") else f"Can't read the date “{r.get('date')}”.")
        at = parse_time(r.get("time")) if r.get("time") is not None else default_time
        if not at:
            issues.append(f"Can't read the time “{r.get('time')}”. Use 09:30 or 9:30 PM.")
        if day and at and datetime.fromisoformat(f"{day}T{at}") <= now_local.replace(tzinfo=None):
            issues.append("That date and time have already passed.")
        chans, unknown = [], []
        for c in split_list(r.get("channels")):
            p = normalize_platform(c)
            p = "x" if p in ("tweet", "x.com") else p
            (chans if p in TITLES else unknown).append(p if p in TITLES else c)
        chans = list(dict.fromkeys(chans))
        if unknown:
            issues.append("Unknown channel: " + ", ".join(unknown) + ". Use " + ", ".join(TITLES.values()) + ".")
        if not chans and not unknown:
            issues.append("Add at least one channel.")
        text = str(r.get("text") or "").strip()
        topic = str(r.get("topic") or "").strip()
        if not text and not topic:
            issues.append("Add the post text, or a topic for the AI to write about.")
        media = split_list(r.get("media"))
        title = str(r.get("title") or "").strip()
        subreddit = str(r.get("subreddit") or "").strip()
        row = {"row": n, "date": day, "time": at, "channels": chans, "title": title, "text": text, "topic": topic,
               "media": media, "firstComment": str(r.get("first_comment") or "").strip(),
               "altText": str(r.get("alt_text") or "").strip(), "subreddit": subreddit}
        issues += problems_for(row, written=bool(text))
        if not text and topic:
            notes.append("The AI writes this post (1 post credit" + ("" if media else " + 1 image credit") + ").")
        if "instagram" in chans and not media and text:
            notes.append("Instagram needs a picture: one is drawn when it posts (1 image credit).")
        row["issues"] = list(dict.fromkeys(issues))
        row["notes"] = notes
        out.append(row)
    return out


def problems_for(row: Dict[str, Any], written: bool = True) -> List[str]:
    """Length, media and Reddit checks per channel, the same ones publishing makes."""
    kinds = [{"url": u, "kind": "video" if re.search(r"\.(mp4|mov)(\?|$)", u, re.I) else "image"} for u in row.get("media") or []]
    shell = SimpleNamespace(title=row.get("title") or default_title(row), copy=row.get("text") or "x", hook="", hashtags=[], cta="",
                            adapt_per_channel=False, linkedin_copy=None, x_copy=None, facebook_copy=None,
                            instagram_copy=None, threads_copy=None, image_url=None, alt_text="", media=kinds,
                            channel_options={"reddit": {"subreddit": row.get("subreddit") or ""}}, channels=row.get("channels") or [])
    out = []
    for i in channel_issues(shell):
        if i["kind"] == "missing_media" or (i["kind"] == "length" and not written and "title" not in i["problem"]):
            continue
        out.append(f"{TITLES.get(i['channel'], i['channel'])}: {i['problem']}")
    return out


def default_title(row: Dict[str, Any]) -> str:
    """A headline for a row without one: the first line of the text (cut at a word), else the topic."""
    first = (row.get("text") or "").strip().split("\n")[0]
    if not first:
        return (row.get("topic") or "Imported post").strip()[:90]
    if len(first) <= 90:
        return first
    cut = first[:90].rsplit(" ", 1)[0]
    return cut.rstrip(" ,;:-") + "…"


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADERS)
    w.writerow(["2026-11-03", "09:30", "linkedin, x", "",
                "Month-end close takes most teams three days. Here is how we got it to one: https://example.com/guide",
                "", "", "", "", ""])
    w.writerow(["2026-11-05", "13:00", "instagram", "", "", "Behind the scenes at our new office",
                "https://example.com/photo1.jpg, https://example.com/photo2.jpg", "", "Team at the new office", ""])
    w.writerow(["2026-11-07", "17:00", "reddit", "What tools do you use for invoicing?",
                "We are a 10-person agency and keep juggling spreadsheets. What works for you?", "", "", "", "", "smallbusiness"])
    return buf.getvalue()
